# -*- coding: utf-8 -*-
"""科研蓝白风格 PPT 基础组件"""
import os
from PIL import Image
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

# ---------- 配色 ----------
NAVY = RGBColor(0x0B, 0x2E, 0x4F)
BLUE = RGBColor(0x1C, 0x5F, 0xA8)
CYAN = RGBColor(0x2E, 0x9B, 0xD6)
TINT = RGBColor(0xEA, 0xF2, 0xFA)
TINT2 = RGBColor(0xF4, 0xF8, 0xFC)
LINE = RGBColor(0xCB, 0xDA, 0xEA)
GRAY = RGBColor(0x5C, 0x68, 0x75)
LGRAY = RGBColor(0x8A, 0x97, 0xA3)
DARK = RGBColor(0x15, 0x22, 0x2E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AMBER = RGBColor(0xC9, 0x7B, 0x1E)

CN = "微软雅黑"
EN = "Arial"

W, H = 13.3333, 7.5
ML, MR = 0.62, 12.7133
CW = MR - ML
FOOTER = "商用自动奶泡系统竞品分析　|　产品开发与市场定位　|　2026-09-18"

_CROP_DIR = "assets_min/_crop"


# ---------- 文本 ----------
def _apply_font(run, size, bold, color, spc=None, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    f.name = EN
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", CN)
    if spc:
        rPr.set("spc", str(int(spc * 100)))


def textbox(slide, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return tf


def para(tf, text, size=11, bold=False, color=GRAY, space_before=0,
         space_after=0, line=1.28, align=None, spc=None, first=False, italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if align is not None:
        p.alignment = align
    p.line_spacing = line
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    r = p.add_run()
    r.text = text
    _apply_font(r, size, bold, color, spc, italic)
    return p


def label(slide, x, y, w, text, size=11, bold=False, color=GRAY, h=0.3,
          align=PP_ALIGN.LEFT, spc=None, line=1.2, anchor=MSO_ANCHOR.TOP, italic=False):
    tf = textbox(slide, x, y, w, h, align=align, anchor=anchor)
    para(tf, text, size, bold, color, align=align, spc=spc, line=line, first=True, italic=italic)
    return tf


def rich(slide, x, y, w, h, parts, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line=1.2):
    """parts: [(text, size, bold, color, spc)]"""
    tf = textbox(slide, x, y, w, h, align=align, anchor=anchor)
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line
    for t, s, b, c, *rest in parts:
        r = p.add_run()
        r.text = t
        _apply_font(r, s, b, c, rest[0] if rest else None)
    return tf


# ---------- 形状 ----------
def rect(slide, x, y, w, h, fill=None, line_color=None, line_w=0.75,
         shape=MSO_SHAPE.RECTANGLE, alpha=None):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        if alpha is not None:
            sf = sh.fill._xPr.find(qn("a:solidFill"))
            clr = sf.find(qn("a:srgbClr"))
            a = etree.SubElement(clr, qn("a:alpha"))
            a.set("val", str(int(alpha * 100000)))
    if line_color is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line_color
        sh.line.width = Pt(line_w)
    sh.text_frame.word_wrap = True
    return sh


def hline(slide, x, y, w, color=LINE, width=0.75):
    from pptx.util import Inches as I
    ln = slide.shapes.add_connector(1, I(x), I(y), I(x + w), I(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(width)
    return ln


def vline(slide, x, y, h, color=LINE, width=0.75):
    from pptx.util import Inches as I
    ln = slide.shapes.add_connector(1, I(x), I(y), I(x), I(y + h))
    ln.line.color.rgb = color
    ln.line.width = Pt(width)
    return ln


# ---------- 图片 ----------
def place_img(slide, path, x, y, w, h, mode="contain", pad=0.0, focus=0.5):
    im = Image.open(path)
    iw, ih = im.size
    ar = iw / ih
    if mode == "contain":
        bw, bh = w - 2 * pad, h - 2 * pad
        if bw / bh > ar:
            dh, dw = bh, bh * ar
        else:
            dw, dh = bw, bw / ar
        return slide.shapes.add_picture(
            path, Inches(x + pad + (bw - dw) / 2), Inches(y + pad + (bh - dh) / 2),
            Inches(dw), Inches(dh))
    target = w / h
    if ar > target:
        nw = int(ih * target)
        left = int((iw - nw) * focus)
        box = (left, 0, left + nw, ih)
    else:
        nh = int(iw / target)
        top = int((ih - nh) * focus)
        box = (0, top, iw, top + nh)
    os.makedirs(_CROP_DIR, exist_ok=True)
    key = "%s_%dx%d_%d.jpg" % (os.path.splitext(os.path.basename(path))[0],
                               int(w * 100), int(h * 100), int(focus * 100))
    out = os.path.join(_CROP_DIR, key)
    if not os.path.exists(out):
        im.convert("RGB").crop(box).resize(
            (min(box[2] - box[0], 1800), int(min(box[2] - box[0], 1800) / target)),
            Image.LANCZOS).save(out, quality=90)
    return slide.shapes.add_picture(path if False else out, Inches(x), Inches(y),
                                    Inches(w), Inches(h))


# ---------- 表格 ----------
_NOSTYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def _cell_border(cell, edge, color, w_pt):
    tcPr = cell._tc.get_or_add_tcPr()
    tag = qn("a:ln" + edge)
    for e in tcPr.findall(tag):
        tcPr.remove(e)
    ln = etree.SubElement(tcPr, tag)
    ln.set("w", str(int(w_pt * 12700)))
    ln.set("cap", "flat")
    ln.set("cmpd", "sng")
    ln.set("algn", "ctr")
    fill = etree.SubElement(ln, qn("a:solidFill"))
    c = etree.SubElement(fill, qn("a:srgbClr"))
    c.set("val", "%02X%02X%02X" % (color[0], color[1], color[2]))
    d = etree.SubElement(ln, qn("a:prstDash"))
    d.set("val", "solid")


def table(slide, x, y, w, col_w, rows, header, body, head_size=10, body_size=9.5,
          head_h=0.40, row_h=0.40, align=None, head_fill=NAVY, zebra=True,
          body_color=DARK, key_col_bold=True):
    """col_w: 各列宽（英寸，和应为 w）；align: 每列对齐"""
    n_rows = len(body) + 1
    gt = slide.shapes.add_table(n_rows, len(col_w), Inches(x), Inches(y),
                                Inches(w), Inches(head_h + row_h * len(body)))
    tbl = gt.table
    tbl._tbl.tblPr.set("firstRow", "0")
    tbl._tbl.tblPr.set("bandRow", "0")
    sid = tbl._tbl.tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tbl._tbl.tblPr, qn("a:tableStyleId"))
    sid.text = _NOSTYLE

    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    tbl.rows[0].height = Inches(head_h)
    for r in range(1, n_rows):
        tbl.rows[r].height = Inches(row_h)

    align = align or [PP_ALIGN.LEFT] * len(col_w)

    def fill_cell(cell, text, size, bold, color, bg, al, borders):
        cell.fill.solid()
        cell.fill.fore_color.rgb = bg
        cell.margin_left = Inches(0.10)
        cell.margin_right = Inches(0.08)
        cell.margin_top = Inches(0.03)
        cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = al
        p.line_spacing = 1.12
        r = p.add_run()
        r.text = text
        _apply_font(r, size, bold, color)
        for e in borders:
            _cell_border(cell, e, LINE, 0.75)

    for c, txt in enumerate(header):
        cell = tbl.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = head_fill
        cell.margin_left = Inches(0.10)
        cell.margin_right = Inches(0.08)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.alignment = align[c]
        r = p.add_run()
        r.text = txt
        _apply_font(r, head_size, True, WHITE, spc=0.4)

    for ri, row in enumerate(body):
        bg = WHITE if (ri % 2 == 0 or not zebra) else TINT2
        for c, txt in enumerate(row):
            bold = key_col_bold and c == 0
            color = NAVY if bold else body_color
            fill_cell(tbl.cell(ri + 1, c), txt, body_size, bold, color, bg,
                      align[c], ["B"])
    return tbl


# ---------- 页面骨架 ----------
def new_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, fill=WHITE)
    return s


def header(slide, num, title, kicker=None):
    rect(slide, ML, 0.50, 0.085, 0.30, fill=CYAN)
    label(slide, ML + 0.21, 0.485, 1.0, num, 13, True, CYAN, spc=1.0)
    label(slide, ML + 0.72, 0.44, 10.6, title, 25, True, NAVY, h=0.5)
    if kicker:
        label(slide, ML, 1.02, CW, kicker, 11, False, GRAY, h=0.3)
    hline(slide, ML, 1.38, CW, LINE, 1.0)


def cover_footer(slide, page):
    hline(slide, ML, 6.99, CW, LINE, 0.75)
    label(slide, ML, 7.08, 9.0, FOOTER, 8.5, False, LGRAY)
    label(slide, MR - 1.0, 7.06, 1.0, "%02d" % page, 10.5, True, NAVY,
          align=PP_ALIGN.RIGHT)


def source_note(slide, text, y=6.68):
    label(slide, ML, y, CW, text, 7.8, False, LGRAY, h=0.24)
