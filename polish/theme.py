# -*- coding: utf-8 -*-
"""视觉系统：配色、字号、栅格与可复用组件。

只负责"外观"层：形状、填充、描边、文字排版、图片定位与裁切。
不生成任何位图，不改动原始图片像素，不改动图表/表格数值。
"""
from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

# ---------------------------------------------------------------- 配色
# 全部沿用原稿已有的蓝白学术色系，只做收敛与统一。
NAVY      = "16416C"   # 主色（原稿标题/分节页）
NAVY_DEEP = "0E2F4F"   # 深一档，用于压角与反白区
NAVY_MID  = "1E5996"   # 原稿母版顶栏蓝
BLUE      = "2A7ABF"   # 亮蓝点缀（原稿分隔线渐变色）
BLUE_TINT = "EDF3F9"   # 浅蓝底板
BLUE_LINE = "CFDEEA"   # 卡片描边
CRIMSON   = "B9282C"   # 原稿母版底栏红 —— 全篇唯一强调红
RED_HI    = "C00000"   # 原稿正文强调红（统一替换 FF0000）
GOLD      = "C8963C"   # 暖色点缀（替换原稿低对比的 FFC000）
GOLD_LT   = "F2C879"   # 反白区强调
GREEN     = "1F7A4D"   # "安全"语义色（替换原稿刺眼的 7EB000）
INK       = "1F2937"   # 正文
INK_MID   = "4B5563"   # 次级正文
INK_SOFT  = "8A94A2"   # 辅助信息
WHITE     = "FFFFFF"
PAPER     = "F7F9FC"

FONT = "微软雅黑"
FONT_EN = "Times New Roman"
FONT_CAL = "华文行楷"      # 校训沿用原稿书法体

# ---------------------------------------------------------------- 栅格
SW, SH = 13.3333, 7.5      # 16:9
M = 0.62                   # 左右页边距
CW = SW - 2 * M            # 内容宽 12.0933
RULE_Y = 1.06              # 标题分隔线
BODY_TOP = 1.28            # 内容区顶
BODY_BOT = 6.52            # 内容区底（常规）
FOOT_Y = 6.93              # 页脚基线
BAR_TOP_H = 0.174          # 母版顶栏高
BAR_BOT_Y = 7.422          # 母版底栏位置


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


def I(inches: float) -> int:
    return Emu(int(round(inches * 914400)))


# ---------------------------------------------------------------- 文字 DSL
class T:
    """一个文字 run。"""

    __slots__ = ("s", "sz", "b", "c", "f", "spc", "i")

    def __init__(self, s, sz=None, b=None, c=None, f=None, spc=None, i=None):
        self.s, self.sz, self.b, self.c, self.f, self.spc, self.i = s, sz, b, c, f, spc, i


class P:
    """一个段落。"""

    __slots__ = ("runs", "al", "sb", "sa", "ls", "lvl", "bullet")

    def __init__(self, runs, al="l", sb=0, sa=0, ls=None, lvl=0, bullet=None):
        if isinstance(runs, (str, T)):
            runs = [runs]
        self.runs = [T(r) if isinstance(r, str) else r for r in runs]
        self.al, self.sb, self.sa, self.ls, self.lvl, self.bullet = al, sb, sa, ls, lvl, bullet


_ALIGN = {
    "l": PP_ALIGN.LEFT,
    "c": PP_ALIGN.CENTER,
    "r": PP_ALIGN.RIGHT,
    "j": PP_ALIGN.JUSTIFY,
    "d": PP_ALIGN.DISTRIBUTE,
}
_ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def _set_run_font(run, name, size, bold, color, spc=None, italic=None):
    f = run.font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    if spc is not None:
        rPr.set("spc", str(int(spc * 100)))
    if name:
        latin = rPr.get_or_add_latin()
        latin.set("typeface", name)
        # a:ea / a:cs 未被 python-pptx 暴露，手动插到 a:latin 之后，
        # 否则中文会回落到主题的 minorFont（等线）。
        prev = latin
        for tag in ("a:ea", "a:cs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = rPr.makeelement(qn(tag), {})
                prev.addnext(el)
            el.set("typeface", name)
            prev = el


def write(tf, paras, *, sz=15, b=False, c=INK, f=FONT, ls=None, al="l"):
    """把段落列表写入 text_frame，第一个段落复用已有空段落。"""
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = _ALIGN[para.al or al]
        if para.sb:
            p.space_before = Pt(para.sb)
        if para.sa:
            p.space_after = Pt(para.sa)
        lsp = para.ls if para.ls is not None else ls
        if lsp is not None:
            p.line_spacing = lsp
        if para.lvl:
            p.level = para.lvl
        # 关掉继承来的项目符号
        pPr = p._pPr if p._pPr is not None else p._p.get_or_add_pPr()
        if pPr.find(qn("a:buNone")) is None and pPr.find(qn("a:buChar")) is None:
            pPr.append(pPr.makeelement(qn("a:buNone"), {}))
        for r in para.runs:
            run = p.add_run()
            run.text = r.s
            _set_run_font(
                run,
                r.f or f,
                r.sz if r.sz is not None else sz,
                r.b if r.b is not None else b,
                r.c or c,
                r.spc,
                r.i,
            )
    return tf


# ---------------------------------------------------------------- 形状基元
def _no_shadow(shape):
    shape.shadow.inherit = False


def soft_shadow(shape, blur=9, dist=2.5, alpha=10, direction=5400000):
    """淡投影，给卡片一点层次。"""
    spPr = shape._element.spPr
    for tag in ("a:effectLst",):
        el = spPr.find(qn(tag))
        if el is not None:
            spPr.remove(el)
    xml = (
        '<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<a:outerShdw blurRad="{int(blur * 12700)}" dist="{int(dist * 12700)}" '
        f'dir="{direction}" rotWithShape="0">'
        f'<a:srgbClr val="1B3A5C"><a:alpha val="{int(alpha * 1000)}"/></a:srgbClr>'
        "</a:outerShdw></a:effectLst>"
    )
    from pptx.oxml import parse_xml

    spPr.append(parse_xml(xml))


def rect(slide, x, y, w, h, *, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE,
         adj=None, shadow=False, alpha=None):
    sp = slide.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    sp.shadow.inherit = False
    if adj is not None:
        try:
            sp.adjustments[0] = adj
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            sf = sp._element.spPr.find(qn("a:solidFill"))
            clr = sf.find(qn("a:srgbClr"))
            a = clr.makeelement(qn("a:alpha"), {"val": str(int(alpha * 1000))})
            clr.append(a)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = rgb(line)
        sp.line.width = Pt(lw)
    if shadow:
        soft_shadow(sp)
    tf = sp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    return sp


def grad(slide, x, y, w, h, c1, c2, *, angle=0, a1=None, a2=None,
         shape=MSO_SHAPE.RECTANGLE, adj=None):
    """双色线性渐变矩形。angle 单位为度，0 = 从左到右。"""
    sp = rect(slide, x, y, w, h, fill=c1, shape=shape, adj=adj)
    f = sp.fill
    f.gradient()
    stops = f.gradient_stops
    stops[0].position = 0.0
    stops[0].color.rgb = rgb(c1)
    stops[1].position = 1.0
    stops[1].color.rgb = rgb(c2)
    for st, av in ((stops[0], a1), (stops[1], a2)):
        if av is not None:
            clr = st.color._xFill.find(qn("a:srgbClr"))
            clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(av * 1000))}))
    f.gradient_angle = angle % 360
    sp.line.fill.background()
    return sp


def textbox(slide, x, y, w, h, paras, *, sz=15, b=False, c=INK, f=FONT, ls=None,
            al="l", anchor="t", pad=0.0):
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = I(pad)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = _ANCHOR[anchor]
    write(tf, paras, sz=sz, b=b, c=c, f=f, ls=ls, al=al)
    return tb


def fill_text(shape, paras, *, sz=15, b=False, c=INK, f=FONT, ls=None, al="l",
              anchor="m", pad=(0.16, 0.16, 0.10, 0.10)):
    """把文字写进一个已存在的形状。pad = (左, 右, 上, 下)。"""
    tf = shape.text_frame
    tf.margin_left, tf.margin_right = I(pad[0]), I(pad[1])
    tf.margin_top, tf.margin_bottom = I(pad[2]), I(pad[3])
    tf.vertical_anchor = _ANCHOR[anchor]
    write(tf, paras, sz=sz, b=b, c=c, f=f, ls=ls, al=al)
    return shape


def line(slide, x1, y1, x2, y2, *, color=NAVY, lw=1.25, dash=None, arrow=None):
    from pptx.util import Emu as _E

    cx = slide.shapes.add_connector(1, I(x1), I(y1), I(x2), I(y2))  # 1 = straight
    cx.line.color.rgb = rgb(color)
    cx.line.width = Pt(lw)
    ln = cx.line._get_or_add_ln()
    if dash:
        el = ln.find(qn("a:prstDash"))
        if el is None:
            el = ln.makeelement(qn("a:prstDash"), {})
            ln.append(el)
        el.set("val", dash)
    if arrow:
        el = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
        ln.append(el)
    return cx


# ---------------------------------------------------------------- 图片
def fit_box(px, py, x, y, w, h):
    """等比缩放放入 (x,y,w,h) 并居中 —— 不裁切，图表/示意图必须走这条。"""
    ar = px / py
    if w / h > ar:
        nh, nw = h, h * ar
    else:
        nw, nh = w, w / ar
    return x + (w - nw) / 2.0, y + (h - nh) / 2.0, nw, nh


def cover_crop(px, py, w, h, *, bias_x=0.5, bias_y=0.5, base=(0.0, 0.0, 0.0, 0.0)):
    """铺满 (w,h) 所需的裁切比例，可叠加在原有裁切之上（仅用于照片）。"""
    bl, br, bt, bb = base
    vis_w = px * (1 - bl - br)
    vis_h = py * (1 - bt - bb)
    ar, target = vis_w / vis_h, w / h
    el = er = et = eb = 0.0
    if ar > target:                       # 图更宽 → 裁左右
        keep = target / ar
        cut = 1 - keep
        el, er = cut * bias_x, cut * (1 - bias_x)
    else:                                 # 图更高 → 裁上下
        keep = ar / target
        cut = 1 - keep
        et, eb = cut * bias_y, cut * (1 - bias_y)
    span_x, span_y = 1 - bl - br, 1 - bt - bb
    return (bl + el * span_x, br + er * span_x, bt + et * span_y, bb + eb * span_y)


def put_pic(slide, path, x, y, w, h, *, crop=None, geom=None, adj=None,
            line_color=None, lw=1.0, shadow=False):
    """放置一张原始图片。python-pptx 按 sha1 复用已有 image part，不会复制媒体。"""
    pic = slide.shapes.add_picture(str(path), I(x), I(y), I(w), I(h))
    if crop:
        pic.crop_left, pic.crop_right, pic.crop_top, pic.crop_bottom = crop
    if geom:
        pg = pic._element.spPr.get_or_add_prstGeom()
        pg.set("prst", geom)
        av = pg.find(qn("a:avLst"))
        if av is None:
            av = pg.makeelement(qn("a:avLst"), {})
            pg.append(av)
        for ch in list(av):
            av.remove(ch)
        if adj is not None:
            av.append(av.makeelement(qn("a:gd"), {"name": "adj", "fmla": "val %d" % int(adj * 100000)}))
    if line_color:
        pic.line.color.rgb = rgb(line_color)
        pic.line.width = Pt(lw)
    if shadow:
        soft_shadow(pic)
    else:
        pic.shadow.inherit = False
    return pic


# ---------------------------------------------------------------- 页面组件
def page_frame(slide):
    """母版同款顶栏 / 底栏，让所有内容页（含空白待补页）视觉一致。"""
    rect(slide, 0, 0, SW, BAR_TOP_H, fill=NAVY)
    rect(slide, 0, BAR_BOT_Y, SW, SH - BAR_BOT_Y, fill=CRIMSON)


_RULE_IMG = None
# 原稿自带的渐变分隔线位图（1977×23，第 13—20 行为渐变，其余为白）。
# 裁到渐变行后横向拉通，作为全篇统一的标题分隔线 —— 复用原图而非另画。
_RULE_CROP = (0.0, 0.0, 13 / 23, 1 - 21 / 23)


def set_rule_image(path):
    global _RULE_IMG
    _RULE_IMG = str(path)


def brand_rule(slide, y, *, x=M, w=None, h=0.035, cap=0.92):
    w = CW * 0.86 if w is None else w
    if _RULE_IMG:
        put_pic(slide, _RULE_IMG, x, y, w, h, crop=_RULE_CROP)
    else:
        grad(slide, x, y, w, h, BLUE, WHITE)
    if cap:
        rect(slide, x, y, cap, h, fill=NAVY)


def head(slide, eyebrow, title_runs, *, sub=None, rule=True, sz=27):
    """统一页眉：小节标签 + 主标题 + 渐变分隔线。"""
    textbox(slide, M, 0.30, CW * 0.7, 0.26,
            [P([T(eyebrow, sz=11.5, b=True, c=INK_SOFT, spc=1.6)])], anchor="t")
    textbox(slide, M, 0.54, CW, 0.50,
            [P(title_runs, ls=1.0)], sz=sz, b=True, c=NAVY, anchor="t")
    if rule:
        brand_rule(slide, RULE_Y)
    if sub:
        textbox(slide, M, RULE_Y + 0.10, CW, 0.28, [P(sub)], sz=13, c=INK_MID)
    return RULE_Y + (0.42 if sub else 0.0)


def foot(slide, idx, total=24, note=None):
    runs = [T("天津科技大学 · 咖啡液萃取与超声杀菌工艺及装备设计研究", sz=9.5, c=INK_SOFT)]
    if note:
        runs.append(T("　|　" + note, sz=9.5, c=INK_SOFT))
    textbox(slide, M, FOOT_Y, CW - 1.50, 0.28, [P(runs)], anchor="t")
    textbox(slide, SW - M - 1.30, FOOT_Y, 1.30, 0.28,
            [P([T("%02d" % idx, sz=12, b=True, c=NAVY), T("  /  %d" % total, sz=9.5, c=INK_SOFT)], al="r")],
            anchor="t")


def card(slide, x, y, w, h, *, fill=WHITE, line_color=BLUE_LINE, radius=0.02, shadow=True):
    sp = rect(slide, x, y, w, h, fill=fill, line=line_color, lw=1.0,
              shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=radius, shadow=shadow)
    return sp


def card_head(slide, x, y, w, h, runs, *, fill=NAVY, al="l", sz=16, radius=0.10):
    """卡片顶部标题条。"""
    sp = rect(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE, adj=radius)
    fill_text(sp, [P(runs, al=al)], sz=sz, b=True, c=WHITE, anchor="m", pad=(0.18, 0.18, 0, 0))
    return sp


def chip(slide, x, y, w, h, runs, *, fill=NAVY, color=WHITE, sz=13, line_color=None, al="c"):
    sp = rect(slide, x, y, w, h, fill=fill, line=line_color,
              shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.5)
    fill_text(sp, [P(runs, al=al)], sz=sz, b=True, c=color, anchor="m", pad=(0.10, 0.10, 0, 0))
    return sp


def takeaway(slide, y, runs, *, h=0.78, x=M, w=CW, sz=16.5, al="c"):
    """结论条：深蓝底反白字 + 左侧红色标识块。"""
    sp = rect(slide, x, y, w, h, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.04)
    rect(slide, x + 0.20, y + 0.14, 0.065, h - 0.28, fill=CRIMSON)
    fill_text(sp, [P(runs, al=al, ls=1.18)], sz=sz, b=True, c=WHITE, anchor="m",
              pad=(0.42, 0.34, 0.04, 0.04))
    return sp


def kpi(slide, x, y, w, h, value, label, *, accent=NAVY, unit=None, note=None):
    """右栏数据卡：大字数值 + 说明。数值全部取自原稿图表/表格，未新增数据。"""
    sp = card(slide, x, y, w, h, fill=WHITE, radius=0.06)
    rect(slide, x, y + 0.14, 0.055, h - 0.28, fill=accent)
    runs = [T(value, sz=27, b=True, c=accent)]
    if unit:
        runs.append(T(" " + unit, sz=13, b=True, c=accent))
    paras = [P(runs, ls=1.0, sa=3), P([T(label, sz=11.5, c=INK_MID)], ls=1.15)]
    if note:
        paras.append(P([T(note, sz=10, c=INK_SOFT)], ls=1.1, sb=2))
    fill_text(sp, paras, anchor="m", pad=(0.24, 0.16, 0.06, 0.06))
    return sp


def placeholder_hint(slide, y, h, text):
    """待补充页的占位提示（用户后续会替换，故用极浅的虚线框）。"""
    sp = rect(slide, M, y, CW, h, fill=PAPER, line=BLUE_LINE, lw=1.0,
              shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.02)
    ln = sp.line._get_or_add_ln()
    el = ln.makeelement(qn("a:prstDash"), {"val": "dash"})
    ln.append(el)
    fill_text(sp, [P([T(text, sz=15, c=INK_SOFT)], al="c")], anchor="m")
    return sp
